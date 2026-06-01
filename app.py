import os
import sys
import math
import random
import traceback
from flask import Flask, jsonify, request, render_template, redirect, url_for

try:
    import openpyxl
    import numpy as np
except ImportError as e:
    print(f"Missing dependency: {e}. Install requirements and run again.")
    sys.exit(1)

app = Flask(__name__, template_folder='templates', static_folder='static')
UPLOAD_FOLDER = os.path.dirname(os.path.abspath(__file__))
DEFAULT_EXCEL = os.path.join(UPLOAD_FOLDER, 'ARCHIVOS.xlsx')
UPLOADED_EXCEL = os.path.join(UPLOAD_FOLDER, 'uploaded_data.xlsx')

def get_active_filepath():
    return UPLOADED_EXCEL if os.path.exists(UPLOADED_EXCEL) else DEFAULT_EXCEL

def normalize_keyword(kw):
    if kw is None: return ""
    s = str(kw).strip()
    if not s: return ""
    return s

def normalize_author(a):
    if a is None: return ""
    return str(a).strip()

def parse_excel_data(filepath):
    if not os.path.exists(filepath):
        return []
    wb = openpyxl.load_workbook(filepath, read_only=True)
    all_articles = []
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        if sheet.max_row <= 1: continue
        header = None
        for r in sheet.iter_rows(max_row=1, values_only=True):
            header = [str(x).strip() if x is not None else "" for x in r]
            break
        if not header or not any(header): continue
        header_lower = [h.lower() for h in header]
        def find_index(possible):
            for name in possible:
                if name.lower() in header_lower:
                    return header_lower.index(name.lower())
            for i,h in enumerate(header_lower):
                for name in possible:
                    if name.lower() in h:
                        return i
            return -1
        idx_key = find_index(['key','id'])
        idx_title = find_index(['title','titulo','document title'])
        idx_author = find_index(['author','authors','autor','autores'])
        idx_journal = find_index(['publication title','journal','source title','revista'])
        idx_tags = find_index(['keywords','palabras clave','tags','automatic tags','manual tags'])
        idx_year = find_index(['year','publication year','año'])
        for row in sheet.iter_rows(min_row=2, values_only=True):
            key = row[idx_key] if idx_key != -1 and idx_key < len(row) else None
            title = row[idx_title] if idx_title != -1 and idx_title < len(row) else None
            if not key and not title: continue
            authors_raw = row[idx_author] if idx_author != -1 and idx_author < len(row) else ""
            journal_raw = row[idx_journal] if idx_journal != -1 and idx_journal < len(row) else ""
            tags_raw = row[idx_tags] if idx_tags != -1 and idx_tags < len(row) else ""
            year_raw = row[idx_year] if idx_year != -1 and idx_year < len(row) else None
            authors = []
            if authors_raw:
                authors = [normalize_author(a) for a in str(authors_raw).split(';') if a.strip()]
            journal = str(journal_raw).strip() if journal_raw else ""
            keywords = []
            if tags_raw:
                sep = ';' if ';' in str(tags_raw) else ','
                for k in [x.strip() for x in str(tags_raw).split(sep) if x.strip()]:
                    keywords.append(normalize_keyword(k))
            try:
                year = int(year_raw) if year_raw else None
            except:
                year = None
            all_articles.append({
                'key': str(key) if key else f'KEY_{random.randint(1000,9999)}',
                'title': str(title) if title else '',
                'authors': authors,
                'journal': journal,
                'keywords': keywords,
                'year': year
            })
    # deduplicate
    seen = set(); out = []
    for a in all_articles:
        uid = a['key']
        if uid in seen: continue
        seen.add(uid); out.append(a)
    return out

def generate_networks(articles):
    kw_freq = {}; auth_freq = {}; journ_freq = {}
    kw_pairs = {}; auth_pairs = {}; journal_keywords = {}
    for art in articles:
        kws = list(set(art.get('keywords', [])))
        for k in kws: kw_freq[k] = kw_freq.get(k,0)+1
        for i in range(len(kws)):
            for j in range(i+1,len(kws)):
                p = tuple(sorted([kws[i], kws[j]])); kw_pairs[p] = kw_pairs.get(p,0)+1
        auths = list(set(art.get('authors', [])))
        for a in auths: auth_freq[a] = auth_freq.get(a,0)+1
        for i in range(len(auths)):
            for j in range(i+1,len(auths)):
                p = tuple(sorted([auths[i], auths[j]])); auth_pairs[p] = auth_pairs.get(p,0)+1
        jn = art.get('journal','')
        if jn:
            journ_freq[jn] = journ_freq.get(jn,0)+1
            journal_keywords.setdefault(jn,set()).update(kws)
    kw_nodes = [{'id':k,'count':v} for k,v in kw_freq.items()]
    kw_edges = [{'source':p[0],'target':p[1],'weight':w} for p,w in kw_pairs.items()]
    auth_nodes = [{'id':k,'count':v} for k,v in auth_freq.items()]
    auth_edges = [{'source':p[0],'target':p[1],'weight':w} for p,w in auth_pairs.items()]
    journ_nodes = [{'id':k,'count':v} for k,v in journ_freq.items()]
    journ_edges = []
    jlist = list(journal_keywords.keys())
    for i in range(len(jlist)):
        for j in range(i+1,len(jlist)):
            shared = journal_keywords[jlist[i]].intersection(journal_keywords[jlist[j]])
            if shared:
                journ_edges.append({'source':jlist[i],'target':jlist[j],'weight':len(shared)})
    return {
        'keywords': {'nodes': sorted(kw_nodes,key=lambda x:-x['count']),'edges': kw_edges},
        'authors': {'nodes': sorted(auth_nodes,key=lambda x:-x['count']),'edges': auth_edges},
        'journals': {'nodes': sorted(journ_nodes,key=lambda x:-x['count']),'edges': journ_edges}
    }

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/api/data')
def api_data():
    try:
        fp = get_active_filepath()
        if not os.path.exists(fp):
            return jsonify({'success': False, 'error': 'Archivo Excel no encontrado'})
        articles = parse_excel_data(fp)
        networks = generate_networks(articles)
        stats = {'articles': len(articles), 'keywords': len(networks['keywords']['nodes']), 'authors': len(networks['authors']['nodes']), 'journals': len(networks['journals']['nodes']), 'source_file': os.path.basename(fp)}
        return jsonify({'success': True, 'stats': stats, 'networks': networks, 'articles': articles})
    except Exception as e:
        traceback.print_exc()
        return jsonify({'success': False, 'error': str(e)})

@app.route('/api/matrix')
def api_matrix():
    try:
        entity = request.args.get('entity','keywords')
        metric = request.args.get('metric','cooccurrence')
        top = int(request.args.get('top','20'))
        fp = get_active_filepath()
        if not os.path.exists(fp):
            return jsonify({'success': False, 'error': 'No Excel loaded'})
        articles = parse_excel_data(fp)
        item_freq = {}; paper_items = []
        for art in articles:
            if entity=='keywords': items = list(set(art.get('keywords',[])))
            elif entity=='authors': items = list(set(art.get('authors',[])))
            else: items = [art.get('journal')] if art.get('journal') else []
            paper_items.append(items)
            for it in items: item_freq[it] = item_freq.get(it,0)+1
        sorted_items = sorted(item_freq.items(), key=lambda x:-x[1])[:top]
        labels = [i for i,_ in sorted_items]
        if not labels: return jsonify({'success': True, 'labels': [], 'matrix': []})
        # build co-occurrence
        all_items = list(item_freq.keys())
        idx = {it:i for i,it in enumerate(all_items)}
        n = len(all_items)
        C = np.zeros((n,n))
        for items in paper_items:
            ids = [idx[it] for it in items if it in idx]
            for a in ids: C[a,a]+=1
            for i in range(len(ids)):
                for j in range(i+1,len(ids)):
                    C[ids[i],ids[j]]+=1; C[ids[j],ids[i]]+=1
        top_idx = [idx[l] for l in labels]
        M = []
        if metric=='cooccurrence':
            for i in top_idx:
                M.append([int(C[i,j]) for j in top_idx])
        elif metric=='cosine':
            for i in top_idx:
                row=[]
                for j in top_idx:
                    denom = math.sqrt(C[i,i]*C[j,j])
                    row.append(round(float(C[i,j]/denom) if denom>0 else 0.0,4))
                M.append(row)
        else:
            for i in top_idx:
                row=[]
                pi = C[i,:]
                for j in top_idx:
                    pj = C[j,:]
                    mi, mj = pi.mean(), pj.mean()
                    di = pi-mi; dj = pj-mj
                    denom = math.sqrt((di**2).sum()*(dj**2).sum())
                    val = float((di*dj).sum()/denom) if denom>0 else 0.0
                    row.append(round(val,4))
                M.append(row)
        return jsonify({'success': True, 'labels': labels, 'matrix': M, 'entity': entity, 'metric': metric})
    except Exception as e:
        traceback.print_exc()
        return jsonify({'success': False, 'error': str(e)})

@app.route('/api/upload', methods=['POST'])
def upload():
    if 'file' not in request.files:
        return redirect(url_for('index'))
    f = request.files['file']
    if f and f.filename.endswith(('.xlsx','.xls')):
        f.save(UPLOADED_EXCEL)
    return redirect(url_for('index'))

@app.route('/api/reset', methods=['POST'])
def reset():
    if os.path.exists(UPLOADED_EXCEL): os.remove(UPLOADED_EXCEL)
    return jsonify({'success': True})

if __name__ == '__main__':
    port = int(os.environ.get('PORT', '5000'))
    app.run(host='0.0.0.0', port=port)
